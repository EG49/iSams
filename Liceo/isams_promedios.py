"""
leer_materias.py  -  PASO 1 (solo lectura, no modifica nada en iSAMS)

Para una promoción (ej. "3BCH") abre, una por una, las materias de
la hoja 'codes' de gradebooks.xlsx por su código (ID) y lee el árbol de estructura (carpetas y
columnas) con el ID de cada fila (el número de [ID123456] en las fórmulas).

Salida: lectura_{promocion}.xlsx
    grado | materia | id_gradebook | fila_id | titulo | tipo | nombre_corto | maximo |
    ponderacion | porcentaje | icono | parent_id | carpeta_padre | texto_fila
  (nombre_corto = columna 'Short Name', ej. 'Actuación clase 20%'; porcentaje = el 20 extraído)

Si una materia falla se reporta y se continúa con la siguiente. El Excel se
guarda tras CADA materia, así que Ctrl+C no pierde lo ya leído.

Uso:  python leer_materias.py 3BCH
"""

import re
import sys
import time

import openpyxl

from datos import (
    LOGIN_URL,
    RUTA_MODULO_GRADEBOOKS,
    RUTA_EDITAR_GRADEBOOK,
    EXCEL_PATH,
    HOJA_CODES,
    PROMOCION_POR_DEFECTO,
    EXCLUIR_MATERIAS,
    ARCHIVO_SALIDA,
)
from funciones_isams import (
    crear_driver,
    login_manual,
    cerrar_alerta_pendiente,
    esperar_filas,
    leer_estructura,
    analizar_fila,
    html_primeras_filas,
    normalizar,
)


def cargar_materias(promocion):
    """Lee la hoja 'codes' (col A = código, col B = Materia con el nombre formal,
    ej. '3BCH-A FÍSICA') y devuelve (id, grado, materia) de las filas cuyo
    grado empieza por la promoción (3BCH -> 3BCH-A, 3BCH-B). El grado es la
    primera palabra del nombre formal."""
    patron = re.compile(r"^" + re.escape(promocion.upper()) + r"(?=[\s\-]|$)")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    hoja = wb[HOJA_CODES]
    materias = []
    for fila in hoja.iter_rows(min_row=2, values_only=True):  # fila 1 = encabezado
        if not fila or fila[0] is None:
            continue
        codigo = str(fila[0]).strip()
        texto = str(fila[1]).strip() if len(fila) > 1 and fila[1] is not None else ""
        if not texto:
            continue
        grado, _, materia = texto.partition(" ")
        if patron.match(grado.upper()):
            materias.append((codigo, grado, materia.strip()))
    return materias


def guardar_excel(ruta, registros):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "estructura"
    ws.append(["grado", "materia", "id_gradebook", "fila_id", "titulo", "tipo", "nombre_corto",
               "maximo", "ponderacion", "porcentaje", "icono", "parent_id", "carpeta_padre", "texto_fila"])
    for r in registros:
        ws.append(r)
    ws.freeze_panes = "A2"
    wb.save(ruta)


def main():
    promocion = (sys.argv[1] if len(sys.argv) > 1 else PROMOCION_POR_DEFECTO).upper()
    excluir = {normalizar(m) for m in EXCLUIR_MATERIAS}
    materias = [m for m in cargar_materias(promocion) if normalizar(m[2]) not in excluir]
    if not materias:
        print(f"No hay materias para '{promocion}' en la hoja '{HOJA_CODES}' de {EXCEL_PATH}.")
        return
    print(f"{promocion}: {len(materias)} materia(s) a leer.")

    driver = crear_driver()
    login_manual(driver)

    # Cargar primero el módulo Gradebooks (carga completa de la SPA). Después, para
    # cada materia solo se cambia el #hash, igual que cuando se navega con clics:
    # abrir directamente la URL con hash en frío no inicializa la vista.
    driver.get(LOGIN_URL.rstrip("/") + "/" + RUTA_MODULO_GRADEBOOKS)
    time.sleep(4)

    registros, reporte = [], []
    diagnostico_guardado = False
    ruta_salida = ARCHIVO_SALIDA.format(promocion=promocion)

    for i, (id_gb, grado, materia) in enumerate(materias, 1):
        print(f"\n[{i}/{len(materias)}] {grado} {materia}  (id={id_gb})")
        try:
            url = LOGIN_URL.rstrip("/") + "/" + RUTA_EDITAR_GRADEBOOK.format(id=id_gb)
            driver.get(url)  # mismo path, solo cambia el hash -> ruta de la SPA, sin recarga
            time.sleep(2)
            cerrar_alerta_pendiente(driver)

            if not esperar_filas(driver, 25):
                # Segundo intento: recarga completa del módulo y de nuevo el hash.
                driver.get(LOGIN_URL.rstrip("/") + "/" + RUTA_MODULO_GRADEBOOKS)
                time.sleep(3)
                driver.get(url)
                time.sleep(2)
                if not esperar_filas(driver, 25):
                    diag = driver.execute_script(
                        "return (document.body ? document.body.innerText : '').slice(0, 200);")
                    raise TimeoutError(
                        f"el árbol de estructura no cargó. URL actual: {driver.current_url} | "
                        f"título: {driver.title!r} | texto: {diag!r}")

            filas = leer_estructura(driver)
            if not filas:
                raise RuntimeError("no se leyó ninguna fila")

            if not diagnostico_guardado:
                with open("diagnostico_filas.html", "w", encoding="utf-8") as fd:
                    fd.write("\n\n<!-- ---- fila ---- -->\n\n".join(html_primeras_filas(driver, 3)))
                diagnostico_guardado = True

            titulo_por_id = {fid: f["titulo"] for fid, f in filas.items()}
            n_carpetas = 0
            for fid, f in filas.items():
                tipo = "carpeta" if f["carpeta"] else "columna"
                n_carpetas += tipo == "carpeta"
                padre = f["parent"] if f["parent"] not in ("", "0") else ""
                corto, maximo, pond, porc = analizar_fila(f)
                registros.append([grado, materia, id_gb, fid, f["titulo"], tipo, corto,
                                  maximo, pond, porc, f.get("icono", ""), padre,
                                  titulo_por_id.get(padre, ""), " | ".join(f.get("textos") or [])])
            print(f"✅ {len(filas)} fila(s): {n_carpetas} carpeta(s), {len(filas) - n_carpetas} columna(s).")
            guardar_excel(ruta_salida, registros)
        except Exception as e:
            msg = f"{grado} {materia} (id={id_gb}): ERROR - {e}"
            print(f"⚠️  {msg}")
            reporte.append(msg)
            cerrar_alerta_pendiente(driver)
            continue

    print("\n" + "=" * 60)
    print(f"Listo. Filas guardadas en '{ruta_salida}': {len(registros)}")
    if reporte:
        print(f"\nMaterias con problema ({len(reporte)}):")
        for r in reporte:
            print(f"  - {r}")
    else:
        print("Todas las materias se leyeron sin errores.")
    input("\nPresiona ENTER para cerrar el navegador...")
    driver.quit()


if __name__ == "__main__":
    main()