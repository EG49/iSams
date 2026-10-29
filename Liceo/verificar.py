"""
verificar.py  -  SOLO LECTURA (no guarda nada en iSAMS)

Revisa que lo hecho esté en orden. Para cada materia (cuantitativa / cualitativa) de la promoción o
curso, en la carpeta CARPETA_A_LEER (TI, TII o TIII):
  1. la carpeta existe y no hay columnas calculadas repetidas;
  2. existen TODAS las columnas calculadas esperadas (COLUMNAS_CUANTITATIVA / COLUMNAS_CUALITATIVA);
  3. los % de los insumos suman 100 (si no, no se pueden calcular las fórmulas esperadas);
  4. (salvo --rapido) cada columna calculada tiene: tipo 'Formula Builder', Allow Sharing =
     COMPARTIR_CALCULADAS y la fórmula IGUAL a la que generaría generar_formulas.py (se comparan con
     nombres: [ID123] -> {NOMBRE}); además avisa si el Short Name no es el esperado.
Materias hija/madre/manual/excluida no se revisan (van en la hoja 'omitidas').

Salida: verificacion_{etiqueta}_{carpeta}.xlsx
    hoja 'resumen'  (una fila por materia: OK / REVISAR)
    hoja 'detalle'  (una fila por chequeo con problema; con --todo también los OK)
    hoja 'omitidas'

Uso:
    python verificar.py 4EGB-A            # un curso    (o una promoción: 4EGB, 3BCH)
    python verificar.py 5EGB 6EGB 7EGB   # varias a la vez (un solo inicio de sesión)
    python verificar.py 4EGB --rapido     # solo estructura (sin abrir cada columna), mucho más rápido
    python verificar.py 4EGB --id 2435    # una sola materia
    python verificar.py 4EGB --todo       # incluye los chequeos OK en 'detalle'
"""

import re
import sys
import time

import openpyxl

from datos import (
    LOGIN_URL, RUTA_MODULO_GRADEBOOKS, CARPETA_A_LEER, EXCLUIR_MATERIAS, COMPARTIR_CALCULADAS,
    TRIMESTRES, COLUMNAS_CUANTITATIVA, COLUMNAS_CUALITATIVA,
)
from funciones_isams import (
    crear_driver, login_manual, cerrar_alerta_pendiente, leer_estructura, analizar_fila,
    filas_de_carpeta, clasificar_materia, normalizar,
)
from generar_formulas import formulas_de_materia, formulas_cualitativa, formulas_cualitativa_mh, _buscar, coincide


def _sin_espacios(texto):
    return re.sub(r"\s+", "", texto or "")


def comparar(filas, categoria, leer=None, cortos=None):
    """filas: [{'fila_id','titulo','porcentaje'}] de la carpeta. leer(titulo) -> (tipo, compartir, formula_con_nombres)
    o None para modo rápido. cortos: {titulo_mayusculas: short name}. Devuelve [(columna, chequeo, estado, esperado, encontrado)]
    donde estado es 'OK' o 'PROBLEMA'."""
    out = []
    n, romano = TRIMESTRES[CARPETA_A_LEER.upper().replace(" ", "")]
    if categoria == "cualitativa_mh":  # 12 columnas por posición; no hay lista de prefijos que revisar
        specs = []
        formulas, error = formulas_cualitativa_mh(filas)
        out.append(("(carpeta)", "12 (o 10) columnas", "OK" if len(filas) in (10, 12) else "PROBLEMA", "12 o 10", str(len(filas))))
    else:
        specs = COLUMNAS_CUALITATIVA if categoria == "cualitativa" else COLUMNAS_CUANTITATIVA

    # 1-2. existencia y repetidos
    presentes = {}
    for prefijo, base, corto in specs:
        coinciden = [f for f in filas if coincide(f["titulo"], prefijo)]
        if not coinciden:
            out.append((f"{base}{n}", "existe", "PROBLEMA", "debe existir", "FALTA la columna"))
        else:
            presentes[prefijo] = coinciden[0]
            if len(coinciden) > 1:
                out.append((coinciden[0]["titulo"], "repetida", "PROBLEMA", "una sola",
                            ", ".join(f["titulo"] for f in coinciden)))
            else:
                out.append((coinciden[0]["titulo"], "existe", "OK", "", ""))
            if cortos is not None:
                esperado = corto.replace("{R}", romano)
                real = cortos.get(coinciden[0]["titulo"].upper(), "")
                if normalizar(real) != normalizar(esperado):
                    out.append((coinciden[0]["titulo"], "short name", "PROBLEMA", esperado, real))

    # 3. fórmulas esperadas
    if categoria != "cualitativa_mh":
        formulas, error = (formulas_cualitativa if categoria == "cualitativa" else formulas_de_materia)(filas)
    if error:
        out.append(("(materia)", "fórmulas esperadas", "PROBLEMA", "se pueden calcular", error))
        return out
    if leer is None:
        return out

    # 4. fórmulas, tipo y compartir en iSAMS
    for f, esperada, _nota, _refs in formulas:
        titulo = f["titulo"]
        try:
            tipo, compartir, real = leer(titulo)
        except Exception as e:
            out.append((titulo, "lectura", "PROBLEMA", "poder leerla", str(e)))
            continue
        out.append((titulo, "tipo de cálculo", "OK" if normalizar(tipo) == normalizar("Formula Builder") else "PROBLEMA",
                    "Formula Builder", tipo))
        out.append((titulo, "compartir", "OK" if normalizar(compartir) == normalizar(COMPARTIR_CALCULADAS) else "PROBLEMA",
                    COMPARTIR_CALCULADAS, compartir))
        if real == esperada:
            out.append((titulo, "fórmula", "OK", esperada, real))
        elif _sin_espacios(real) == _sin_espacios(esperada):
            out.append((titulo, "fórmula", "OK", esperada, real + "   (solo difieren espacios)"))
        else:
            out.append((titulo, "fórmula", "PROBLEMA", esperada, real))
    return out


def main():
    args = sys.argv[1:]
    rapido = "--rapido" in args
    todo = "--todo" in args
    solo_id = args[args.index("--id") + 1] if "--id" in args else None
    pos = [a for i, a in enumerate(args) if not a.startswith("--") and not (i > 0 and args[i - 1] == "--id")]
    if not pos:
        print("Indica el curso o promoción, ej.: python verificar.py 4EGB-A")
        return
    etiquetas = [p.upper() for p in pos]
    etiqueta = "_".join(etiquetas)
    if CARPETA_A_LEER.upper().replace(" ", "") not in TRIMESTRES:
        sys.exit(f"CARPETA_A_LEER='{CARPETA_A_LEER}' no está en TRIMESTRES (datos.py).")

    from leer_materias import cargar_materias
    from escribir_formulas import abrir_materia
    excluir = {normalizar(m) for m in EXCLUIR_MATERIAS}
    materias = [m for e in etiquetas for m in cargar_materias(e) if normalizar(m[2]) not in excluir]
    if solo_id:
        materias = [m for m in materias if m[0] == solo_id]
    if not materias:
        print(f"No hay materias para '{etiqueta}'.")
        return
    print(f"Verificando {etiqueta}, carpeta '{CARPETA_A_LEER}', {len(materias)} materia(s)" + (" (modo rápido)." if rapido else "."))

    out = openpyxl.Workbook()
    wr = out.active
    wr.title = "resumen"
    wr.append(["grado", "materia", "id_gradebook", "categoria", "estado", "problemas"])
    wd = out.create_sheet("detalle")
    wd.append(["grado", "materia", "id_gradebook", "columna", "chequeo", "estado", "esperado", "encontrado"])
    wo = out.create_sheet("omitidas")
    wo.append(["grado", "materia", "id_gradebook", "categoria"])
    ruta = f"verificacion_{etiqueta}_{CARPETA_A_LEER.upper()}.xlsx"

    driver = crear_driver()
    login_manual(driver)
    driver.get(LOGIN_URL.rstrip("/") + "/" + RUTA_MODULO_GRADEBOOKS)
    time.sleep(4)

    malas = []
    try:
        for i, (id_gb, grado, materia) in enumerate(materias, 1):
            categoria = clasificar_materia(grado, materia)
            if categoria not in ("cuantitativa", "cualitativa", "cualitativa_mh"):
                wo.append([grado, materia, id_gb, categoria])
                continue
            print(f"[{i}/{len(materias)}] {grado} {materia} ({categoria})")
            try:
                abrir_materia(driver, id_gb)
                estructura = leer_estructura(driver)
                hijas = filas_de_carpeta(estructura, CARPETA_A_LEER)
                if hijas is None:
                    raise RuntimeError(f"no tiene la carpeta '{CARPETA_A_LEER}'")
                filas, cortos = [], {}
                for _fid, f in hijas:
                    corto, _maximo, _p, porc = analizar_fila(f)
                    filas.append({"fila_id": str(_fid), "titulo": f["titulo"], "porcentaje": porc})
                    cortos[f["titulo"].upper()] = corto
                leer = None
                if not rapido:
                    from leer_formulas import leer_una
                    estado_materia = {"id_a_nombre": None}
                    def leer(titulo, _e=estado_materia):
                        tipo, compartir, _f, con_nombres = leer_una(driver, titulo, _e)
                        return tipo, compartir, con_nombres
                resultados = comparar(filas, categoria, leer, cortos)
            except Exception as e:
                resultados = [("(materia)", "abrir/leer", "PROBLEMA", "", str(e))]
                cerrar_alerta_pendiente(driver)

            problemas = [r for r in resultados if r[2] != "OK"]
            for col, chequeo, estado, esperado, encontrado in resultados:
                if estado != "OK" or todo:
                    wd.append([grado, materia, id_gb, col, chequeo, estado, esperado, encontrado])
            if problemas:
                malas.append((grado, materia, problemas))
                print(f"   ⚠️  {len(problemas)} problema(s): " + "; ".join(f"{p[0]} ({p[1]})" for p in problemas[:4]))
            wr.append([grado, materia, id_gb, categoria, "REVISAR" if problemas else "OK", len(problemas)])
            out.save(ruta)
    except KeyboardInterrupt:
        print("\nInterrumpido; lo revisado hasta ahora está en", ruta)
    finally:
        out.save(ruta)

    total = wr.max_row - 1
    print("\n" + "=" * 60)
    print(f"Listo. {total - len(malas)} materia(s) OK, {len(malas)} para revisar. Detalle en '{ruta}'.")
    for grado, materia, probs in malas:
        print(f"  - {grado} {materia}: " + "; ".join(f"{p[0]} ({p[1]})" for p in probs[:3]) + (" ..." if len(probs) > 3 else ""))
    input("\nPresiona ENTER para cerrar el navegador...")
    driver.quit()


if __name__ == "__main__":
    main()