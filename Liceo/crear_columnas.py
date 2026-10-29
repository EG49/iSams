"""
crear_columnas.py  -  crea las columnas calculadas que FALTAN en la carpeta (trimestre) y les pone su fórmula.

Por defecto NO abre el navegador: solo escribe plan_columnas_{promocion}.xlsx con lo que crearía
(hoja 'crear' con nombre, short name, fórmula; hoja 'existentes'; hoja 'problemas') para revisarlo.
Con --ejecutar abre iSAMS (login manual) y las crea en orden, una por una:
   menú 'Add Calculated Column' -> Name, Short Name, Position (carpeta), Formula Builder, Rounding
   (given decimal, REDONDEO_DECIMALES), Allow Sharing -> Next -> fórmula -> Next ... -> Finish.
Las columnas que ya existen en la carpeta se saltan. Se necesita la lectura de la carpeta:
   1) pon CARPETA_A_LEER = "TII" en datos.py   2) python leer_materias.py 3BCH
   3) python crear_columnas.py 3BCH            (plan)      4) python crear_columnas.py 3BCH --ejecutar

Uso: python crear_columnas.py 2EGB [--ejecutar] [--max N] [--id <id_gradebook>]
"""

import sys
import time

import openpyxl

from datos import (
    LOGIN_URL, RUTA_MODULO_GRADEBOOKS, PROMOCION_POR_DEFECTO, ARCHIVO_SALIDA, CARPETA_A_LEER,
    TRIMESTRES, REDONDEO_TIPO, REDONDEO_DECIMALES, ARCHIVO_PLAN_COLUMNAS, ARCHIVO_RESULTADO_CREACION,
    COLUMNAS_CUANTITATIVA, COLUMNAS_CUALITATIVA,
)
from funciones_isams import (
    clasificar_materia, crear_driver, login_manual, cerrar_alerta_pendiente,
    abrir_asistente_nueva_columna, rellenar_propiedades_nueva_columna, terminar_asistente_creacion,
    abrir_pestana_columns, mapear_columnas, escribir_formula,
)
from generar_formulas import formulas_de_materia, formulas_cualitativa, _buscar, coincide


def planificar(grado, materia, filas, categoria):
    """Devuelve (a_crear, existentes, error). a_crear = [dict(nombre, nombre_corto, formula)] en orden de creación."""
    n, romano = TRIMESTRES[CARPETA_A_LEER.upper().replace(" ", "")]
    specs = COLUMNAS_CUALITATIVA if categoria == "cualitativa" else COLUMNAS_CUANTITATIVA
    existentes = [f["titulo"] for f in filas for p, _, _ in specs if coincide(f["titulo"], p)]
    sinteticas = []
    nuevas = {}
    for prefijo, base, corto in specs:
        if _buscar(filas, prefijo) is None:
            nombre = f"{base}{n}"
            fila = {"fila_id": "", "titulo": nombre, "porcentaje": None}
            sinteticas.append(fila)
            nuevas[nombre.upper()] = corto.replace("{R}", romano)
    if not sinteticas:
        return [], existentes, None
    todas = filas + sinteticas
    formulas, error = (formulas_cualitativa if categoria == "cualitativa" else formulas_de_materia)(todas)
    if error:
        return [], existentes, error
    por_titulo = {f["titulo"].upper(): formula for f, formula, _, _ in formulas}
    a_crear = []
    for fila in sinteticas:  # orden de las specs = orden de creación
        formula = por_titulo.get(fila["titulo"].upper())
        if formula is None:
            return [], existentes, f"no se generó fórmula para {fila['titulo']}"
        a_crear.append({"nombre": fila["titulo"], "nombre_corto": nuevas[fila["titulo"].upper()], "formula": formula})
    return a_crear, existentes, None


def crear_una(driver, col, carpeta, mapa_estado):
    from escribir_formulas import resolver_formula
    abrir_asistente_nueva_columna(driver)
    rellenar_propiedades_nueva_columna(driver, col["nombre"], col["nombre_corto"], carpeta, REDONDEO_TIPO, REDONDEO_DECIMALES)
    abrir_pestana_columns(driver)
    mapa = mapear_columnas(driver)  # incluye las columnas creadas antes en esta misma materia
    formula = resolver_formula(col["formula"], mapa)
    escribir_formula(driver, formula)
    terminar_asistente_creacion(driver, col["nombre_corto"])
    return formula


def main():
    args = sys.argv[1:]
    ejecutar = "--ejecutar" in args
    maximo = int(args[args.index("--max") + 1]) if "--max" in args else None
    solo_id = args[args.index("--id") + 1] if "--id" in args else None
    pos = [a for i, a in enumerate(args) if not a.startswith("--") and not (i > 0 and args[i - 1] in ("--max", "--id"))]
    promocion = (pos[0] if pos else PROMOCION_POR_DEFECTO).upper()
    if CARPETA_A_LEER.upper().replace(" ", "") not in TRIMESTRES:
        sys.exit(f"CARPETA_A_LEER='{CARPETA_A_LEER}' no está en TRIMESTRES (datos.py).")

    wb = openpyxl.load_workbook(ARCHIVO_SALIDA.format(promocion=promocion), data_only=True)
    grupos = {}
    for fila in wb[CARPETA_A_LEER].iter_rows(min_row=2, values_only=True):
        if not fila or fila[3] is None:
            continue
        grado, materia, id_gb, fila_id, titulo, porc = (list(fila) + [None] * 6)[:6]
        porc = float(porc) if porc not in (None, "") else None
        grupos.setdefault((grado, materia, str(id_gb)), []).append(
            {"fila_id": str(fila_id), "titulo": str(titulo or "").strip(), "porcentaje": porc})
    if solo_id:
        grupos = {k: v for k, v in grupos.items() if k[2] == solo_id}

    out = openpyxl.Workbook()
    wc = out.active
    wc.title = "crear"
    wc.append(["grado", "materia", "id_gradebook", "carpeta", "orden", "nombre", "nombre_corto", "formula"])
    we = out.create_sheet("existentes")
    we.append(["grado", "materia", "id_gradebook", "columnas que ya existen"])
    wp = out.create_sheet("problemas")
    wp.append(["grado", "materia", "id_gradebook", "problema"])
    plan = {}
    for (grado, materia, id_gb), filas in grupos.items():
        categoria = clasificar_materia(grado, materia)
        if categoria not in ("cuantitativa", "cualitativa"):
            continue
        a_crear, existentes, error = planificar(grado, materia, filas, categoria)
        if error:
            wp.append([grado, materia, id_gb, f"{error}  |  columnas que hay en {CARPETA_A_LEER}: "
                       + (", ".join(f["titulo"] for f in filas) or "(ninguna)")])
            continue
        if existentes:
            we.append([grado, materia, id_gb, ", ".join(existentes)])
        if a_crear:
            plan[(grado, materia, id_gb)] = a_crear
            for i, c in enumerate(a_crear, 1):
                wc.append([grado, materia, id_gb, CARPETA_A_LEER, i, c["nombre"], c["nombre_corto"], c["formula"]])
    ruta_plan = ARCHIVO_PLAN_COLUMNAS.format(promocion=promocion)
    out.save(ruta_plan)
    total = sum(len(v) for v in plan.values())
    print(f"{promocion} / {CARPETA_A_LEER}: {total} columna(s) por crear en {len(plan)} materia(s); "
          f"{wp.max_row - 1} materia(s) con problemas -> '{ruta_plan}'")
    if not ejecutar:
        print("Es solo el plan. Revísalo y corre de nuevo con --ejecutar para crearlas en iSAMS.")
        return
    if maximo:
        plan = dict(list(plan.items())[:maximo])

    from escribir_formulas import abrir_materia
    ruta_res = ARCHIVO_RESULTADO_CREACION.format(promocion=promocion)
    res = openpyxl.Workbook()
    wr = res.active
    wr.title = "resultado"
    wr.append(["grado", "materia", "id_gradebook", "columna", "estado", "detalle"])
    driver = crear_driver()
    login_manual(driver)
    driver.get(LOGIN_URL.rstrip("/") + "/" + RUTA_MODULO_GRADEBOOKS)
    time.sleep(4)
    errores = []
    try:
        for i, ((grado, materia, id_gb), cols) in enumerate(plan.items(), 1):
            print(f"\n[{i}/{len(plan)}] {grado} {materia} (id={id_gb})")
            try:
                abrir_materia(driver, id_gb)
            except Exception as e:
                errores.append(f"{grado} {materia}: no abrió - {e}")
                print("⚠️ ", errores[-1])
                wr.append([grado, materia, id_gb, "", "error", f"no abrió: {e}"])
                res.save(ruta_res)
                cerrar_alerta_pendiente(driver)
                continue
            for col in cols:
                try:
                    crear_una(driver, col, CARPETA_A_LEER, {})
                    print(f"✅ {col['nombre']} ({col['nombre_corto']})")
                    wr.append([grado, materia, id_gb, col["nombre"], "ok", ""])
                except Exception as e:
                    msg = f"{grado} {materia} / {col['nombre']}: ERROR - {e}"
                    print("⚠️ ", msg)
                    errores.append(msg)
                    wr.append([grado, materia, id_gb, col["nombre"], "error", str(e)])
                    cerrar_alerta_pendiente(driver)
                    try:
                        abrir_materia(driver, id_gb)  # deja la materia limpia; las siguientes se omitirían por dependencia
                    except Exception:
                        pass
                    break  # las columnas siguientes dependen de esta
                res.save(ruta_res)
            res.save(ruta_res)
    except KeyboardInterrupt:
        print("\nInterrumpido.")
    finally:
        res.save(ruta_res)
    print("\n" + "=" * 60)
    print(f"Listo. {len(errores)} error(es). Detalle en '{ruta_res}'.")
    for e in errores:
        print("  -", e)
    input("\nPresiona ENTER para cerrar el navegador...")
    driver.quit()


if __name__ == "__main__":
    main()