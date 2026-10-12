"""
escribir_formulas.py  -  PASO 3 (ESCRIBE en iSAMS)

Lee formulas_{promocion}.xlsx (generado por generar_formulas.py y revisado por ti)
y, materia por materia, abre cada columna calculada y escribe su fórmula:

    ubicar la fila por su NOMBRE dentro de la carpeta (ej. 'SUMAINST1' en 'TI') -> lápiz ->
    Step 1: Formula Builder + Not Shared -> Next -> pestaña 'Columns' (se leen los IDs
    reales, los de [ID77590]) -> se reemplaza cada {NOMBRE} de la plantilla por [ID...] ->
    escribir la fórmula -> Next, Next, Next -> Finish

Los ids del árbol de estructura NO se usan (son números de fila que iSAMS reasigna al guardar).
Si un nombre no aparece en 'Columns' o aparece repetido, esa fórmula NO se escribe y se reporta.

Regla de siempre: si una fórmula falla, se reporta y se continúa con la siguiente.
resultado_{promocion}.xlsx se guarda tras CADA fórmula (ok / error). Con --reanudar se saltan
las que ya quedaron 'ok'.

Uso:
    python escribir_formulas.py 3BCH                 # todas las materias
    python escribir_formulas.py 3BCH --max 1         # solo las primeras N materias (para probar)
    python escribir_formulas.py 3BCH --id 2285       # solo la materia con ese ID de gradebook
    python escribir_formulas.py 3BCH --reanudar      # continúa lo que quedó pendiente/con error
"""

import re
import sys
import time

import openpyxl

from datos import (
    LOGIN_URL,
    RUTA_MODULO_GRADEBOOKS,
    RUTA_EDITAR_GRADEBOOK,
    PROMOCION_POR_DEFECTO,
    ARCHIVO_FORMULAS,
    ARCHIVO_RESULTADO,
)
from funciones_isams import (
    crear_driver,
    login_manual,
    cerrar_alerta_pendiente,
    esperar_filas,
    buscar_elementos_con_scroll,
    buscar_id_carpeta,
    buscar_boton_editar_columna,
    clic_seguro,
    xpath_literal,
    normalizar,
    configurar_propiedades_calculo,
    abrir_pestana_columns,
    mapear_columnas,
    escribir_formula,
    guardar_wizard_calculo,
)


def cargar_formulas(promocion):
    """{(grado, materia, id_gb): [dict(orden, carpeta, titulo, formula, referencias)]}"""
    wb = openpyxl.load_workbook(ARCHIVO_FORMULAS.format(promocion=promocion), data_only=True)
    grupos = {}
    for fila in wb["formulas"].iter_rows(min_row=2, values_only=True):
        if not fila or fila[0] is None:
            continue
        grado, materia, id_gb, carpeta, orden, titulo, formula, _nota, refs = (list(fila) + [None] * 9)[:9]
        grupos.setdefault((grado, materia, str(id_gb)), []).append({
            "orden": int(orden), "carpeta": str(carpeta), "titulo": str(titulo),
            "formula": str(formula), "referencias": [r for r in str(refs or "").split(";") if r],
        })
    for lista in grupos.values():
        lista.sort(key=lambda d: d["orden"])
    return grupos


def cargar_resultados(ruta):
    """(set de (id_gb, carpeta, fila) ya 'ok', lista de filas del resultado anterior)."""
    try:
        wb = openpyxl.load_workbook(ruta, data_only=True)
    except Exception:
        return set(), []
    filas, ok = [], set()
    for f in wb.active.iter_rows(min_row=2, values_only=True):
        if not f or f[0] is None:
            continue
        filas.append(list(f))
        if f[6] == "ok":
            ok.add((str(f[2]), str(f[3]), str(f[4])))
    return ok, filas


def guardar_resultados(ruta, filas):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "resultado"
    ws.append(["grado", "materia", "id_gradebook", "carpeta", "fila", "formula", "estado", "detalle"])
    for r in filas:
        ws.append(r)
    ws.freeze_panes = "A2"
    wb.save(ruta)


def abrir_materia(driver, id_gb):
    url = LOGIN_URL.rstrip("/") + "/" + RUTA_EDITAR_GRADEBOOK.format(id=id_gb)
    driver.get(url)  # mismo path del módulo: solo cambia el #hash (ruta de la SPA)
    time.sleep(2)
    cerrar_alerta_pendiente(driver)
    if not esperar_filas(driver, 25):
        driver.get(LOGIN_URL.rstrip("/") + "/" + RUTA_MODULO_GRADEBOOKS)
        time.sleep(3)
        driver.get(url)
        time.sleep(2)
        if not esperar_filas(driver, 25):
            raise TimeoutError(f"el árbol no cargó. URL actual: {driver.current_url}")


def ubicar_fila(driver, titulo, carpeta, tiempo_espera=45):
    """Fila-columna con ese título DENTRO de la carpeta (el mismo nombre puede repetirse en
    otra carpeta). El id de la carpeta se busca cada vez porque iSAMS reasigna ids al guardar."""
    carpeta_id = buscar_id_carpeta(driver, carpeta, tiempo_espera)
    filas = buscar_elementos_con_scroll(
        driver,
        f"//div[contains(@class,'row')][@title={xpath_literal(titulo)}][@parent='{carpeta_id}']",
        tiempo_espera=tiempo_espera,
    )
    if not filas:
        raise TimeoutError(f"no se encontró la columna '{titulo}' dentro de la carpeta '{carpeta}'")
    return filas[0]


def resolver_formula(plantilla, mapa):
    """Reemplaza cada {NOMBRE} por [ID<id>]. Falla si un nombre no existe o está repetido."""
    def reemplazo(m):
        nombre = m.group(1)
        ids = mapa.get(normalizar(nombre), [])
        if not ids:
            raise RuntimeError(f"'{nombre}' no aparece en la pestaña Columns")
        if len(ids) > 1:
            raise RuntimeError(f"'{nombre}' está repetido en el gradebook (ids {ids}); no se sabe cuál usar")
        return f"[ID{ids[0]}]"
    return re.sub(r"\{([^{}]+)\}", reemplazo, plantilla)


def escribir_una(driver, f, estado_materia):
    """estado_materia['mapa'] = {nombre_normalizado: [ids]} (se lee una vez por materia)."""
    fila = ubicar_fila(driver, f["titulo"], f["carpeta"])
    clic_seguro(driver, buscar_boton_editar_columna(driver, fila))
    time.sleep(1.5)
    configurar_propiedades_calculo(driver)

    if estado_materia.get("mapa") is None:
        abrir_pestana_columns(driver)
        estado_materia["mapa"] = mapear_columnas(driver)
        print(f"   (leídos los IDs de {sum(len(v) for v in estado_materia['mapa'].values())} columna(s))")

    formula = resolver_formula(f["formula"], estado_materia["mapa"])
    escribir_formula(driver, formula)
    guardar_wizard_calculo(driver)
    return formula


def main():
    args = sys.argv[1:]
    reanudar = "--reanudar" in args
    maximo = int(args[args.index("--max") + 1]) if "--max" in args else None
    solo_id = args[args.index("--id") + 1] if "--id" in args else None
    posicionales = [a for i, a in enumerate(args)
                    if not a.startswith("--") and not (i > 0 and args[i - 1] in ("--max", "--id"))]
    promocion = (posicionales[0] if posicionales else PROMOCION_POR_DEFECTO).upper()

    grupos = cargar_formulas(promocion)
    if solo_id:
        grupos = {k: v for k, v in grupos.items() if k[2] == solo_id}
        if not grupos:
            print(f"El ID {solo_id} no está en formulas_{promocion}.xlsx (revisa la hoja 'problemas').")
            return
    if maximo:
        grupos = dict(list(grupos.items())[:maximo])
    total = sum(len(v) for v in grupos.values())
    print(f"{promocion}: {len(grupos)} materia(s), {total} fórmula(s) a escribir.")

    ruta_res = ARCHIVO_RESULTADO.format(promocion=promocion)
    ya_ok, resultados = cargar_resultados(ruta_res) if reanudar else (set(), [])
    if reanudar:
        resultados = [r for r in resultados if r[6] == "ok"]  # los errores se reintentan y se re-registran
        print(f"Reanudando: {len(ya_ok)} fórmula(s) ya estaban ok.")

    driver = crear_driver()
    login_manual(driver)
    driver.get(LOGIN_URL.rstrip("/") + "/" + RUTA_MODULO_GRADEBOOKS)
    time.sleep(4)

    errores = []
    try:
        for i, ((grado, materia, id_gb), lista) in enumerate(grupos.items(), 1):
            print(f"\n[{i}/{len(grupos)}] {grado} {materia} (id={id_gb})")
            pendientes = [f for f in lista if (id_gb, f["carpeta"], f["titulo"]) not in ya_ok]
            if not pendientes:
                print("   todo ok de corridas anteriores, se salta.")
                continue
            try:
                abrir_materia(driver, id_gb)
            except Exception as e:
                msg = f"{grado} {materia}: ERROR al abrir - {e}"
                print(f"⚠️  {msg}")
                errores.append(msg)
                for f in pendientes:
                    resultados.append([grado, materia, id_gb, f["carpeta"], f["titulo"], f["formula"], "error", f"no abrió: {e}"])
                guardar_resultados(ruta_res, resultados)
                cerrar_alerta_pendiente(driver)
                continue

            estado_materia = {"mapa": None}
            for f in pendientes:
                try:
                    formula = escribir_una(driver, f, estado_materia)
                    print(f"✅ {f['titulo']}: {formula}")
                    resultados.append([grado, materia, id_gb, f["carpeta"], f["titulo"], formula, "ok", ""])
                except Exception as e:
                    msg = f"{grado} {materia} / {f['titulo']}: ERROR - {e}"
                    print(f"⚠️  {msg}")
                    errores.append(msg)
                    resultados.append([grado, materia, id_gb, f["carpeta"], f["titulo"], f["formula"], "error", str(e)])
                    cerrar_alerta_pendiente(driver)
                    # Si quedó el wizard abierto, recargamos la materia para dejarla limpia.
                    try:
                        abrir_materia(driver, id_gb)
                    except Exception:
                        pass
                guardar_resultados(ruta_res, resultados)
    except KeyboardInterrupt:
        print("\nInterrumpido; lo escrito hasta ahora está en", ruta_res)
    finally:
        guardar_resultados(ruta_res, resultados)

    ok = sum(1 for r in resultados if r[6] == "ok")
    print("\n" + "=" * 60)
    print(f"Listo. {ok} fórmula(s) ok, {len(errores)} con error. Detalle en '{ruta_res}'.")
    for e in errores:
        print(f"  - {e}")
    input("\nPresiona ENTER para cerrar el navegador...")
    driver.quit()


if __name__ == "__main__":
    main()