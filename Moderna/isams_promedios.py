"""
Automatización iSAMS: agrega columnas compartidas (Promedios/Quimestres/etc.)
a cada gradebook "Promedios {grado}", según los grados listados en un Excel.

Configuración (rutas, hoja del Excel, columnas a marcar) vive en datos.py.

Flujo general:
1. Abre la URL de login.
2. El usuario inicia sesión manualmente en el navegador.
3. Al presionar ENTER, el script:
   - Lee los grados y los códigos de materia 
   - Por cada grado: busca su gradebook "Promedios {grado}", abre
     "Add Shared Columns", selecciona la materia correspondiente a ese
     grado exacto, marca las columnas configuradas.
4. Al final imprime un resumen de todo lo que no se pudo encontrar.

NOTA SOBRE SHADOW DOM:
Esta plataforma anida componentes con shadow DOM "abierto", a veces varios
niveles seguidos, pero en muchos casos el contenido real es DOM ligero
normal (el shadow solo contiene el "chrome" del componente, vía <slot>).
Dos funciones auxiliares cubren ambos casos:

- elemento_profundo(driver, selector): busca un elemento en cualquier parte
  del documento, atravesando cualquier cantidad de shadow roots anidados.
- dentro_de_shadow(driver, host, selector): dado un host ya localizado,
  busca DENTRO de su propio shadow root (un solo nivel).

Cuando la duda es real, usa capturar_ruta.js (pegar en la consola del
navegador y hacer clic en el elemento) para confirmar la ruta exacta antes
de adivinar selectores.
"""

import re
import time
import unicodedata

import openpyxl
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.ui import WebDriverWait, Select
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import UnexpectedAlertPresentException
from selenium.webdriver.chrome.service import Service

from datos import (
    CHROME_PATH,
    DRIVER_PATH,
    LOGIN_URL,
    EXCEL_PATH,
    HOJA_GRADES,
    HOJA_CODES,
    COLUMNAS_A_MARCAR,
)

# --- Funciones auxiliares para shadow DOM ---

_DEEP_QUERY_JS = """
function recolectar(root, selector, resultados) {
    resultados.push(...root.querySelectorAll(selector));
    const todos = root.querySelectorAll('*');
    for (const nodo of todos) {
        if (nodo.shadowRoot) {
            recolectar(nodo.shadowRoot, selector, resultados);
        }
    }
}
const resultados = [];
recolectar(document, arguments[0], resultados);
// Esta plataforma reutiliza ids genericos (form, input, position, wizard...)
// en muchos componentes, y los dialogos de Kendo suelen quedar ocultos en el
// DOM al cerrarse en vez de eliminarse. Por eso, entre varias coincidencias,
// preferimos la que este realmente visible en pantalla (la actual), no la
// primera que aparezca en el orden del documento (que podria ser un dialogo
// viejo ya cerrado).
for (const el of resultados) {
    if (el.offsetParent !== null || el.getClientRects().length > 0) {
        return el;
    }
}
return resultados.length > 0 ? resultados[0] : null;
"""


def elemento_profundo(driver, selector, tiempo_espera=20):
    """Busca un elemento en cualquier parte del documento, atravesando
    cualquier cantidad de shadow roots anidados."""
    fin = time.time() + tiempo_espera
    while time.time() < fin:
        el = driver.execute_script(_DEEP_QUERY_JS, selector)
        if el is not None:
            return el
        time.sleep(0.3)
    raise TimeoutError(f"No se encontró el elemento (búsqueda profunda): {selector}")


def elemento_profundo_por_texto(driver, selector_contenedor, texto_exacto, tiempo_espera=20):
    """Como elemento_profundo, pero filtra por el texto visible exacto del
    elemento (útil cuando varios elementos comparten la misma clase), y
    entre los que coincidan prefiere el que esté realmente visible en
    pantalla (por los mismos diálogos viejos ocultos que motivan
    elemento_profundo)."""
    script = """
        function recolectar(root, selector, resultados) {
            resultados.push(...root.querySelectorAll(selector));
            const todos = root.querySelectorAll('*');
            for (const nodo of todos) {
                if (nodo.shadowRoot) {
                    recolectar(nodo.shadowRoot, selector, resultados);
                }
            }
        }
        const candidatos = [];
        recolectar(document, arguments[0], candidatos);
        const coincidencias = candidatos.filter(
            (c) => (c.textContent || '').trim() === arguments[1]
        );
        for (const el of coincidencias) {
            if (el.offsetParent !== null || el.getClientRects().length > 0) {
                return el;
            }
        }
        return coincidencias.length > 0 ? coincidencias[0] : null;
    """
    fin = time.time() + tiempo_espera
    while time.time() < fin:
        el = driver.execute_script(script, selector_contenedor, texto_exacto)
        if el is not None:
            return el
        time.sleep(0.3)
    raise TimeoutError(f"No se encontró '{texto_exacto}' con selector: {selector_contenedor}")


def dentro_de_shadow(driver, host_element, selector_interno, tiempo_espera=20):
    """Dado un elemento host ya localizado, busca un elemento dentro de su
    propio shadow root (un solo nivel, no recursivo)."""
    fin = time.time() + tiempo_espera
    while time.time() < fin:
        el = driver.execute_script(
            "return arguments[0].shadowRoot ? arguments[0].shadowRoot.querySelector(arguments[1]) : null;",
            host_element, selector_interno,
        )
        if el is not None:
            return el
        time.sleep(0.3)
    raise TimeoutError(f"No se encontró dentro del shadow root: {selector_interno}")


def dentro_de_shadow_todos(driver, host_element, selector_interno):
    """Como dentro_de_shadow, pero devuelve TODOS los elementos que matchean
    dentro del shadow root del host (un solo nivel), no solo el primero."""
    return driver.execute_script(
        "return arguments[0].shadowRoot ? Array.from(arguments[0].shadowRoot.querySelectorAll(arguments[1])) : [];",
        host_element, selector_interno,
    )


def clic_seguro(driver, elemento):
    """Intenta un clic normal; si falla, dispara el clic vía JavaScript."""
    try:
        elemento.click()
    except Exception:
        driver.execute_script("arguments[0].click();", elemento)


def clic_unico_para_alerta(driver, elemento):
    """Da UN solo clic, sin reintentar con un segundo clic de respaldo (a
    diferencia de clic_seguro). Se usa para elementos que disparan un
    alert() nativo síncrono: si reintentáramos con un clic de JS tras la
    excepción del primero, dispararíamos el alert DOS veces, dejando uno
    huérfano bloqueando toda la sesión de Selenium."""
    try:
        elemento.click()
    except UnexpectedAlertPresentException:
        pass  # esperado: el alert ya se abrió durante este mismo clic


def cerrar_alerta_pendiente(driver):
    """Si por algún motivo quedó un alert nativo abierto (de un paso
    anterior que falló a mitad de camino), lo cierra antes de seguir, para
    que no bloquee todos los comandos de Selenium siguientes."""
    try:
        alerta = driver.switch_to.alert
        alerta.accept()
        return True
    except Exception:
        return False


def seleccionar_opcion_con_flechas(driver, select_element, valor, max_intentos=1500):
    """Hace clic en la primera <option> del select (para enfocarlo) y luego
    navega con la flecha ↓ una opción a la vez, comprobando el 'value'
    actual en cada paso, hasta llegar a la opción buscada."""
    # Esperar a que el select esté realmente interactuable antes de
    # intentar nada: al abrir "Add Shared Columns" para la 2da/3ra materia
    # seguida (sin recargar página), el modal a veces tarda un instante en
    # asentarse tras su animación de apertura.
    WebDriverWait(driver, 20).until(lambda d: select_element.is_displayed() and select_element.is_enabled())

    opciones = select_element.find_elements(By.TAG_NAME, "option")
    if not opciones:
        raise RuntimeError("El select no tiene ninguna opción cargada todavía.")

    clic_seguro(driver, opciones[0])
    valor_actual = driver.execute_script("return arguments[0].value;", select_element)

    intentos = 0
    while valor_actual != valor and intentos < max_intentos:
        try:
            select_element.send_keys(Keys.ARROW_DOWN)
        except Exception:
            # El select puede quedar momentáneamente no-interactuable
            # (re-render, animación); esperamos un poco y reintentamos el
            # mismo paso en vez de propagar el error.
            time.sleep(0.5)
            continue
        valor_actual = driver.execute_script("return arguments[0].value;", select_element)
        intentos += 1

    if valor_actual != valor:
        raise TimeoutError(f"No se llegó a la opción con value={valor} tras {intentos} flechas.")
    return valor_actual


def clic_boton_wizard(driver, wizard_selector, texto_boton, tiempo_espera=20):
    """Dentro del shadow root del wizard, busca entre los isams-button de
    #wizardButtons el que tenga ese texto visible (insensible a mayúsculas)
    y le da clic."""
    wizard = elemento_profundo(driver, wizard_selector, tiempo_espera)
    fin = time.time() + tiempo_espera
    while time.time() < fin:
        botones_host = dentro_de_shadow_todos(driver, wizard, "#wizardButtons isams-button")
        for host in botones_host:
            try:
                btn = driver.execute_script(
                    "return arguments[0].shadowRoot ? arguments[0].shadowRoot.querySelector('button.button') : null;",
                    host,
                )
                if btn is None:
                    continue
                texto = (btn.get_attribute("textContent") or "").strip()
            except Exception:
                continue
            if texto.lower() == texto_boton.lower():
                # Durante la transición entre pasos, el botón puede quedar
                # deshabilitado unos instantes mientras el contenido del
                # siguiente paso termina de cargar de forma asíncrona.
                try:
                    if not (btn.is_displayed() and btn.is_enabled()):
                        continue
                except Exception:
                    continue
                clic_seguro(driver, btn)
                return btn
        time.sleep(0.3)
    raise TimeoutError(f"No se encontró el botón del wizard con texto '{texto_boton}'")


def xpath_literal(texto):
    """Literal seguro para XPath, incluso si el texto tiene comillas."""
    if "'" not in texto:
        return f"'{texto}'"
    if '"' not in texto:
        return f'"{texto}"'
    partes = texto.split("'")
    return "concat(" + ", \"'\", ".join(f"'{p}'" for p in partes) + ")"


def seleccionar_columna_por_nombre(driver, nombre_columna, tiempo_espera=20):
    """En la grilla de columnas compartidas (DOM plano, sin shadow), busca
    la fila <tr> con un <span title="{nombre_columna}"> y marca su checkbox."""
    xpath = f"//tr[.//span[@title={xpath_literal(nombre_columna)}]]//input[@type='checkbox']"
    checkbox = WebDriverWait(driver, tiempo_espera).until(
        EC.element_to_be_clickable((By.XPATH, xpath))
    )
    clic_seguro(driver, checkbox)
    return checkbox


def buscar_fila_por_nombre(driver, nombre_buscado):
    """Recorre las filas de la tabla de resultados y devuelve la que tenga
    ese nombre exacto (vive dentro del shadow DOM del <isams-link>)."""
    filas = driver.find_elements(By.CSS_SELECTOR, "tr[data-uid]")
    for fila in filas:
        try:
            link = fila.find_element(By.TAG_NAME, "isams-link")
            a_tag = link.shadow_root.find_element(By.CSS_SELECTOR, "a")
            texto = (a_tag.get_attribute("textContent") or "").strip()
        except Exception:
            continue
        if texto == nombre_buscado:
            return fila
    return None


# --- Utilidades de datos (Excel / matching de materia) ---

def normalizar(texto):
    """Mayúsculas, sin tildes, espacios colapsados, para comparar textos con
    formato inconsistente (MAYÚSCULAS/Capitalizado, con/sin tildes, etc.)."""
    texto = (texto or "").strip().upper()
    texto = "".join(c for c in unicodedata.normalize("NFKD", texto) if not unicodedata.combining(c))
    texto = " ".join(texto.split())
    return texto


def cargar_datos_excel():
    """Lee la hoja 'grades' (columna A, sin encabezado, un grado por fila)
    y la hoja 'codes' (columnas Codigo | Materia, con encabezado en la
    fila 1) del Excel configurado en datos.py."""
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

    hoja_grades = wb[HOJA_GRADES]
    grados = []
    for (valor,) in hoja_grades.iter_rows(min_col=1, max_col=1, values_only=True):
        if valor is None or str(valor).strip() == "":
            continue
        grados.append(str(valor).strip())

    hoja_codes = wb[HOJA_CODES]
    codigos = []
    for fila in hoja_codes.iter_rows(min_row=2, values_only=True):  # fila 1 = encabezado
        if not fila or fila[0] is None:
            continue
        codigo = str(fila[0]).strip()
        materia_texto = str(fila[1]).strip() if len(fila) > 1 and fila[1] is not None else ""
        codigos.append((codigo, materia_texto))

    return grados, codigos


def buscar_icono_en_fila(driver, row_element, clase_icono, tiempo_espera=20):
    """Dado un <div class='row'> del árbol de estructura, ubica un ícono
    específico (por su clase, ej. 'icon-16-column-edit' o
    'icon-16-branch-element') dentro del shadow DOM de algún isams-icon
    hijo, sin importar en qué posición esté entre los demás iconos de la
    fila, y devuelve el host (isams-icon) al que hay que darle clic."""
    script = """
        function buscarProfundoDesde(root, selector) {
            let el = root.querySelector(selector);
            if (el) return el;
            const todos = root.querySelectorAll('*');
            for (const nodo of todos) {
                if (nodo.shadowRoot) {
                    el = buscarProfundoDesde(nodo.shadowRoot, selector);
                    if (el) return el;
                }
            }
            return null;
        }
        const icono = buscarProfundoDesde(arguments[0], 'i.' + arguments[1]);
        if (!icono) return null;
        const raiz = icono.getRootNode();
        return raiz.host || icono;
    """
    fin = time.time() + tiempo_espera
    while time.time() < fin:
        host = driver.execute_script(script, row_element, clase_icono)
        if host is not None:
            return host
        time.sleep(0.3)
    raise TimeoutError(f"No se encontró el ícono '{clase_icono}' en la fila.")


def buscar_boton_editar_columna(driver, row_element, tiempo_espera=20):
    """Ubica el ícono de editar (lápiz) de una fila del árbol de estructura."""
    return buscar_icono_en_fila(driver, row_element, "icon-16-column-edit", tiempo_espera)


def extraer_nombre_materia(materia_texto, grado):
    """Quita el prefijo del grado (ej. '5EGB-A') del texto de materia tal
    como viene en la hoja 'codes' (ej. '5EGB-A - Estudios Sociales' o
    '5EGB-A ESTUDIOS SOCIALES'), preservando mayúsculas/tildes originales
    del resto, para usarlo como nombre de carpeta/Position."""
    texto = (materia_texto or "").strip()
    patron = re.compile(r"^" + re.escape(grado.strip()) + r"\s*[-]?\s*", re.IGNORECASE)
    return patron.sub("", texto).strip()


def clic_next_wizard_verificado(driver, wizard_selector, tiempo_espera=20):
    """Como clic_boton_wizard con texto 'Next', pero además verifica que el
    indicador 'Step X of Y' del wizard realmente haya cambiado antes de
    continuar (por si el botón queda deshabilitado momentáneamente mientras
    el siguiente paso carga contenido de forma asíncrona, ej. una grilla de
    'Report Mappings')."""
    wizard = elemento_profundo(driver, wizard_selector, tiempo_espera)
    paging = dentro_de_shadow(driver, wizard, "#wizardPaging", tiempo_espera)
    texto_antes = (paging.get_attribute("textContent") or "").strip()

    fin = time.time() + tiempo_espera
    while time.time() < fin:
        try:
            clic_boton_wizard(driver, wizard_selector, "Next", tiempo_espera=5)
        except TimeoutError:
            pass  # el botón puede seguir deshabilitado; reintentamos abajo

        sub_fin = time.time() + 5
        while time.time() < sub_fin:
            wizard_actual = elemento_profundo(driver, wizard_selector, tiempo_espera)
            paging_actual = dentro_de_shadow(driver, wizard_actual, "#wizardPaging", tiempo_espera)
            texto_actual = (paging_actual.get_attribute("textContent") or "").strip()
            if texto_actual != texto_antes:
                return texto_actual
            time.sleep(0.3)
        # No cambió de paso todavía: volvemos a intentar el clic en Next

    raise TimeoutError(f"El wizard no avanzó de '{texto_antes}' tras reintentar 'Next' varias veces.")


def buscar_ligero_desde(driver, elemento_raiz, selector, tiempo_espera=20):
    """Busca dentro del DOM ligero (sin cruzar shadow) de un elemento ya
    localizado, usando querySelector normal escopado a ese elemento. Útil
    para evitar ambigüedad con ids genéricos reutilizados en otras partes
    de la página, una vez que ya tenemos el contenedor correcto (ej. el
    wizard actualmente abierto)."""
    fin = time.time() + tiempo_espera
    while time.time() < fin:
        el = driver.execute_script("return arguments[0].querySelector(arguments[1]);", elemento_raiz, selector)
        if el is not None:
            return el
        time.sleep(0.3)
    raise TimeoutError(f"No se encontró (DOM ligero) dentro del elemento: {selector}")


_ENCONTRAR_CONTENEDOR_SCROLLEABLE_JS = """
function encontrarContenedorScrolleable(el) {
    let nodo = el;
    while (nodo && nodo !== document.body) {
        if (nodo.scrollHeight > nodo.clientHeight + 5) {
            return nodo;
        }
        nodo = nodo.parentElement || (nodo.getRootNode() && nodo.getRootNode().host);
    }
    return document.scrollingElement || document.documentElement;
}
const ancla = document.querySelector('div.row');
if (!ancla) return null;
return encontrarContenedorScrolleable(ancla);
"""


def buscar_elementos_con_scroll(driver, xpath, filtro=None, tiempo_espera=45):
    """Busca elementos por XPath en el árbol de estructura de columnas
    (markbooks-structure), que resultó ser VIRTUALIZADO: solo las filas
    visibles en pantalla existen en el DOM. Si no aparecen coincidencias,
    hace scroll incremental dentro del contenedor scrolleable más cercano
    (ubicado dinámicamente a partir de cualquier fila conocida) y reintenta,
    hasta encontrar algo o llegar al final del scroll (y de vuelta al
    inicio, por si la fila buscada quedó arriba). Devuelve una lista
    (puede estar vacía si no se encuentra nada)."""
    fin = time.time() + tiempo_espera
    ultimo_scroll_top = -1
    reinicio_probado = False

    while time.time() < fin:
        candidatos = driver.find_elements(By.XPATH, xpath)
        if filtro:
            candidatos = [c for c in candidatos if filtro(c)]
        if candidatos:
            return candidatos

        contenedor = driver.execute_script(_ENCONTRAR_CONTENEDOR_SCROLLEABLE_JS)
        if contenedor is None:
            time.sleep(0.3)
            continue

        nuevo_scroll_top = driver.execute_script(
            "arguments[0].scrollTop += arguments[0].clientHeight * 0.8; return arguments[0].scrollTop;",
            contenedor,
        )
        if nuevo_scroll_top == ultimo_scroll_top:
            if reinicio_probado:
                break  # ya dimos toda la vuelta (fin -> inicio) sin encontrar nada
            # Llegamos al final del scroll sin encontrar nada; probamos
            # también desde el inicio por si la fila buscada quedó arriba.
            driver.execute_script("arguments[0].scrollTop = 0;", contenedor)
            reinicio_probado = True
            ultimo_scroll_top = 0
            time.sleep(0.3)
            continue
        ultimo_scroll_top = nuevo_scroll_top
        time.sleep(0.3)

    return []


def buscar_id_carpeta(driver, nombre_carpeta, tiempo_espera=45):
    """Devuelve el id numérico de la fila-carpeta (tiene el atributo 'group')
    con ese título. Compara de forma insensible a mayúsculas/tildes, y
    scrollea el árbol (virtualizado) si hace falta hasta encontrarla."""
    objetivo_norm = normalizar(nombre_carpeta)
    filas = buscar_elementos_con_scroll(
        driver,
        "//div[contains(@class,'row')][@group]",
        filtro=lambda f: normalizar(f.get_attribute("title") or "") == objetivo_norm,
        tiempo_espera=tiempo_espera,
    )
    if not filas:
        raise TimeoutError(f"No se encontró la carpeta '{nombre_carpeta}' en el árbol de estructura.")
    try:
        return int(filas[0].get_attribute("id"))
    except (TypeError, ValueError):
        raise RuntimeError(f"La carpeta '{nombre_carpeta}' no tiene un id numérico válido.")


def buscar_fila_estructura_mas_reciente(driver, nombre_columna, tiempo_espera=45):
    """Busca la fila SUELTA (parent="0") con ese título exacto en el árbol
    de estructura - es decir, la columna recién agregada por 'Add Shared
    Columns' que todavía no se movió a ninguna carpeta. Scrollea el árbol
    (virtualizado) si hace falta hasta encontrarla.

    IMPORTANTE: nunca cae a "cualquier fila con ese título" como respaldo.
    Si iSAMS todavía no terminó de crear la columna del lado del servidor
    (tarda más a medida que se acumulan columnas en la misma sesión), no
    existe ninguna fila parent="0" todavía; en ese caso seguimos esperando
    hasta el timeout. Agarrar una fila ya asignada a otra materia (aunque
    tenga el mismo nombre) sería mover la columna equivocada."""
    filas = buscar_elementos_con_scroll(
        driver,
        f"//div[contains(@class,'row')][@title={xpath_literal(nombre_columna)}][@parent='0']",
        tiempo_espera=tiempo_espera,
    )
    if not filas:
        raise TimeoutError(
            f"'{nombre_columna}' no apareció como columna suelta (parent=0) tras {tiempo_espera}s. "
            f"Probablemente iSAMS todavía no terminó de crearla del lado del servidor."
        )

    def id_numerico(fila):
        try:
            return int(fila.get_attribute("id"))
        except (TypeError, ValueError):
            return -1
    return max(filas, key=id_numerico)


def editar_columna_asignar_posicion(driver, nombre_columna, nombre_carpeta, materia_texto, tiempo_espera=45):
    """Abre el editor de una columna compartida ya añadida (por su nombre) y
    le asigna la carpeta/materia correcta en el campo 'Position', avanzando
    el resto del wizard (Next -> Next -> Next -> Finish). Luego verifica DOS
    cosas, sin corregir nada automáticamente si fallan (solo se reporta):
    1) que el 'parent' de la fila coincida con el id de la carpeta destino;
    2) que el ícono de origen ('Shared From', un alert nativo) diga que la
       columna viene exactamente de materia_texto (ej. '5EGB-D - Matemáticas'),
       para descartar que se haya compartido la materia equivocada."""
    cerrar_alerta_pendiente(driver)  # por si quedó algo abierto de un intento anterior
    fila = buscar_fila_estructura_mas_reciente(driver, nombre_columna, tiempo_espera)

    boton_editar = buscar_boton_editar_columna(driver, fila, tiempo_espera)
    clic_seguro(driver, boton_editar)
    time.sleep(1.5)

    # Escopamos la búsqueda del formulario al wizard ACTUAL (ya localizado y
    # visible) en vez de buscar "isams-data-form#form" en todo el documento,
    # que podría chocar con un diálogo viejo oculto con el mismo id genérico.
    wizard = elemento_profundo(driver, "isams-wizard#wizard", tiempo_espera)
    form = buscar_ligero_desde(driver, wizard, "isams-data-form#form", tiempo_espera)
    dropdown_position = dentro_de_shadow(driver, form, "isams-dropdown#position", tiempo_espera)
    select_position = dentro_de_shadow(driver, dropdown_position, "select#input", tiempo_espera)

    try:
        Select(select_position).select_by_visible_text(nombre_carpeta)
    except Exception:
        objetivo_norm = normalizar(nombre_carpeta)
        encontrado = False
        for opcion in select_position.find_elements(By.TAG_NAME, "option"):
            if normalizar(opcion.text) == objetivo_norm:
                clic_seguro(driver, opcion)
                encontrado = True
                break
        if not encontrado:
            raise RuntimeError(f"No se encontró la posición/carpeta '{nombre_carpeta}' en el select.")

    time.sleep(0.5)
    clic_next_wizard_verificado(driver, "isams-wizard#wizard")
    time.sleep(0.5)
    clic_next_wizard_verificado(driver, "isams-wizard#wizard")
    time.sleep(0.5)
    clic_next_wizard_verificado(driver, "isams-wizard#wizard")
    time.sleep(0.5)
    clic_boton_wizard(driver, "isams-wizard#wizard", "Finish")
    time.sleep(1.5)

    # --- Verificación 1: iSAMS re-renderiza TODO el árbol de estructura al
    # guardar (reasignando ids nuevos a todas las filas), así que cualquier
    # id que hubiéramos guardado antes de editar queda obsoleto de inmediato.
    # En vez de eso, volvemos a buscar la fila por su título DESPUÉS del
    # Finish, y confirmamos que alguna coincidencia tenga el 'parent'
    # correcto (id de la carpeta destino).
    carpeta_id = buscar_id_carpeta(driver, nombre_carpeta, tiempo_espera)
    candidatas_correctas = buscar_elementos_con_scroll(
        driver,
        f"//div[contains(@class,'row')][@title={xpath_literal(nombre_columna)}]",
        filtro=lambda c: c.get_attribute("parent") == str(carpeta_id),
        tiempo_espera=tiempo_espera,
    )
    fila_verificar = candidatas_correctas[0] if candidatas_correctas else None

    if fila_verificar is None:
        raise RuntimeError(
            f"'{nombre_columna}' no quedó dentro de '{nombre_carpeta}' (id de carpeta: {carpeta_id})"
        )

    # --- Verificación 2: el ícono "Shared From" abre un alert NATIVO del
    # navegador con el gradebook de origen real (ej. '5EGB-D - Matemáticas').
    # Comparamos contra materia_texto tal como viene de 'codes'; si no
    # coincide, reportamos el problema pero NO movemos ni corregimos nada
    # automáticamente (podría ser un desfase real que hay que revisar a mano).
    icono_origen = buscar_icono_en_fila(driver, fila_verificar, "icon-16-branch-element", tiempo_espera)
    clic_unico_para_alerta(driver, icono_origen)

    WebDriverWait(driver, tiempo_espera).until(EC.alert_is_present())
    alerta = driver.switch_to.alert
    texto_alerta = alerta.text
    alerta.accept()

    match = re.search(r"&quot;(.*?)&quot;", texto_alerta)
    texto_origen = match.group(1) if match else texto_alerta

    if normalizar(texto_origen) != normalizar(materia_texto):
        raise RuntimeError(
            f"'{nombre_columna}' dice venir de '{texto_origen}' "
            f"(se esperaba '{materia_texto}') - revisar manualmente"
        )


def buscar_materias_por_grado(codigos, grado):
    """Dentro de la lista (codigo, materia_texto) de la hoja 'codes', devuelve
    TODAS las filas cuyo texto empiece EXACTAMENTE por el grado dado (para
    que '5EGB-A' no matchee con '5EGB-AB' ni '5EGB-D'). Cada una es una
    materia distinta de ese grado (Estudios Sociales, Science, etc.)."""
    grado_norm = normalizar(grado)
    patron_grado = re.compile(r"^" + re.escape(grado_norm) + r"(?=[\s\-]|$)")

    resultado = []
    for codigo, materia_texto in codigos:
        texto_norm = normalizar(materia_texto)
        if patron_grado.match(texto_norm):
            resultado.append((codigo, materia_texto))
    return resultado


# --- Pasos de alto nivel ---

def ir_a_lista_gradebooks(driver, wait):
    """Navega al módulo Gradebooks, busca 'promedio' y filtra por
    'Show All Gradebooks', dejando la tabla de resultados lista."""
    modulo_url = LOGIN_URL.rstrip("/") + "/main/areas/markbooks/index.html"
    driver.get(modulo_url)
    wait.until(EC.presence_of_element_located((By.ID, "filterGradebooks")))

    contenedor_busqueda = elemento_profundo(driver, "isams-textbox#filterGradebooks")
    campo_texto = dentro_de_shadow(driver, contenedor_busqueda, "input[placeholder='Quick Search Gradebooks...']")
    campo_texto.clear()
    campo_texto.send_keys("promedio")

    dropdown_contenedor = elemento_profundo(driver, "isams-dropdown#userSelection")
    select_el = dentro_de_shadow(driver, dropdown_contenedor, "select#input")
    Select(select_el).select_by_visible_text("Show All Gradebooks")

    time.sleep(2)


def procesar_grado(driver, wait, grado, codigos):
    """Ejecuta el flujo completo de 'Add Shared Columns' para TODAS las
    materias de un grado (cada una encontrada en la hoja 'codes'). Devuelve
    una lista de strings describiendo lo que no se pudo encontrar (vacía si
    todo salió bien)."""
    faltantes = []
    print(f"\n=== Procesando grado: {grado} ===")

    ir_a_lista_gradebooks(driver, wait)

    nombre_fila = f"Promedios {grado}"
    fila_objetivo = buscar_fila_por_nombre(driver, nombre_fila)
    if fila_objetivo is None:
        mensaje = f"{grado}: gradebook '{nombre_fila}' no encontrado en la lista"
        print(f"⚠️  {mensaje}")
        return [mensaje]

    icono = fila_objetivo.find_element(By.CSS_SELECTOR, "isams-icon[title='View Gradebook']")
    clic_seguro(driver, icono)

    WebDriverWait(driver, 20).until(lambda d: re.search(r"/markbooks/\w+/(\d+)", d.current_url))
    match = re.search(r"/markbooks/\w+/(\d+)", driver.current_url)
    gradebook_id = match.group(1)

    edit_url = LOGIN_URL.rstrip("/") + f"/main/areas/markbooks/index.html#/markbooks/edit/{gradebook_id}"
    driver.get(edit_url)
    time.sleep(2)

    materias = buscar_materias_por_grado(codigos, grado)
    if not materias:
        mensaje = f"{grado}: no se encontró ninguna materia en la hoja '{HOJA_CODES}' para este grado"
        print(f"⚠️  {mensaje}")
        return [mensaje]

    print(f"{len(materias)} materia(s) encontradas para '{grado}': " + ", ".join(m[1] for m in materias))

    for codigo, materia_texto in materias:
        print(f"\n-- Materia: {materia_texto} (código {codigo}) --")
        try:
            item_shared_columns = elemento_profundo_por_texto(driver, "div.options-menu-item", "Add Shared Columns")
            clic_seguro(driver, item_shared_columns)
            time.sleep(1.5)

            listbox = elemento_profundo(driver, "isams-listbox#existingMarkbookColumns")
            select_columnas = dentro_de_shadow(driver, listbox, "select#input")
            seleccionar_opcion_con_flechas(driver, select_columnas, codigo)

            time.sleep(1.5)  # dar tiempo a que la grilla de abajo se actualice

            columnas_agregadas = []
            for nombre_columna in COLUMNAS_A_MARCAR:
                try:
                    seleccionar_columna_por_nombre(driver, nombre_columna, tiempo_espera=5)
                    print(f"  ✅ {nombre_columna}")
                    columnas_agregadas.append(nombre_columna)
                except Exception:
                    mensaje = f"{grado} / {materia_texto}: columna '{nombre_columna}' no encontrada"
                    print(f"  ⚠️  {mensaje}")
                    faltantes.append(mensaje)

            time.sleep(1)
            clic_next_wizard_verificado(driver, "isams-wizard#existingAssessmentWizard")
            time.sleep(1)
            clic_boton_wizard(driver, "isams-wizard#existingAssessmentWizard", "Finish")
            # Dar tiempo a que iSAMS termine de crear las columnas del lado
            # del servidor antes de empezar a moverlas una por una (esto se
            # vuelve más lento a medida que se acumulan columnas en la misma
            # sesión). buscar_fila_estructura_mas_reciente igual reintenta
            # por su cuenta si hace falta más tiempo todavía.
            time.sleep(3)
            print(f"  ✅ Materia '{materia_texto}' agregada.")

            # --- Mover cada columna recién añadida a la carpeta de su materia
            # (usando el campo "Position" del editor, en vez de drag & drop) ---
            nombre_carpeta = extraer_nombre_materia(materia_texto, grado)
            for nombre_columna in columnas_agregadas:
                try:
                    editar_columna_asignar_posicion(driver, nombre_columna, nombre_carpeta, materia_texto)
                    print(f"  📁 '{nombre_columna}' movida a la carpeta '{nombre_carpeta}'.")
                except Exception as e:
                    mensaje = f"{grado} / {materia_texto}: no se pudo mover '{nombre_columna}' a '{nombre_carpeta}' - {e}"
                    print(f"  ⚠️  {mensaje}")
                    faltantes.append(mensaje)
                    cerrar_alerta_pendiente(driver)  # evitar que un alert huérfano arrastre las siguientes columnas
        except Exception as e:
            mensaje = f"{grado} / {materia_texto}: ERROR al procesar - {e}"
            print(f"⚠️  {mensaje}")
            faltantes.append(mensaje)
            cerrar_alerta_pendiente(driver)

    print(f"✅ Grado '{grado}' completado ({len(materias)} materia(s)).")
    return faltantes


# --- Programa principal ---

def main():
    print(f"Leyendo '{EXCEL_PATH}'...")
    grados, codigos = cargar_datos_excel()
    print(f"{len(grados)} grados cargados de la hoja '{HOJA_GRADES}'.")
    print(f"{len(codigos)} códigos de materia cargados de la hoja '{HOJA_CODES}'.")

    options = webdriver.ChromeOptions()
    if CHROME_PATH:
        options.binary_location = CHROME_PATH

    if DRIVER_PATH:
        service = Service(DRIVER_PATH)
        driver = webdriver.Chrome(service=service, options=options)
    else:
        driver = webdriver.Chrome(options=options)

    driver.maximize_window()
    driver.get(LOGIN_URL)

    input("\n>>> Inicia sesión manualmente en la ventana del navegador.\n"
          ">>> Cuando ya estés dentro de la plataforma, vuelve aquí y presiona ENTER...\n")

    wait = WebDriverWait(driver, 20)
    reporte_faltantes = []

    for grado in grados:
        try:
            faltantes = procesar_grado(driver, wait, grado, codigos)
            reporte_faltantes.extend(faltantes)
        except Exception as e:
            mensaje = f"{grado}: ERROR inesperado - {e}"
            print(f"❌ {mensaje}")
            reporte_faltantes.append(mensaje)
            cerrar_alerta_pendiente(driver)

    print("\n===== RESUMEN =====")
    print(f"Grados procesados: {len(grados)}")
    if reporte_faltantes:
        print(f"Elementos no encontrados ({len(reporte_faltantes)}):")
        for item in reporte_faltantes:
            print(f"  - {item}")
    else:
        print("Todo se procesó sin faltantes. ✅")

    input("\nPresiona ENTER para cerrar el navegador...")
    driver.quit()


if __name__ == "__main__":
    main()
    